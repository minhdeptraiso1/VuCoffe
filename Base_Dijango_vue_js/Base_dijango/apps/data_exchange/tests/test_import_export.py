import io

import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from openpyxl import load_workbook

from apps.accounts.roles import ROLE_ADMIN, assign_role
from apps.accounts.tests.factories import UserFactory
from apps.catalog.models import Category
from apps.data_exchange.services import build_template, parse_workbook
from apps.inventory.models import Unit


def category_file() -> bytes:
    workbook = load_workbook(io.BytesIO(build_template("categories")))
    workbook["Data"].append(["new-coffee", "Cà phê mới", "", 10, True])
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


@pytest.mark.django_db
def test_import_preview_confirm_is_idempotent(api_client):
    admin = UserFactory(is_superuser=True, is_staff=True)
    assign_role(user=admin, role=ROLE_ADMIN)
    api_client.force_authenticate(admin)
    uploaded = SimpleUploadedFile(
        "categories.xlsx",
        category_file(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    preview = api_client.post(
        "/api/v1/imports/",
        {"module": "categories", "mode": "upsert", "file": uploaded},
        format="multipart",
    )
    assert preview.status_code == 201
    assert preview.data["status"] == "ready"
    url = f"/api/v1/imports/{preview.data['id']}/confirm/"
    confirmed = api_client.post(url, {"idempotency_key": "import-1"}, format="json")
    assert confirmed.status_code == 200
    repeated = api_client.post(url, {"idempotency_key": "import-1"}, format="json")
    assert repeated.status_code == 200
    assert Category.objects.filter(code="new-coffee").count() == 1


@pytest.mark.django_db
def test_export_escapes_formula_injection(api_client):
    admin = UserFactory(is_superuser=True, is_staff=True)
    assign_role(user=admin, role=ROLE_ADMIN)
    Category.objects.create(code="safe", name='=HYPERLINK("bad")')
    api_client.force_authenticate(admin)
    response = api_client.get("/api/v1/exports/categories/")
    assert response.status_code == 200
    workbook = load_workbook(io.BytesIO(response.content), data_only=False)
    assert workbook["Data"]["B3"].value.startswith("'")


@pytest.mark.django_db
def test_export_reuses_the_import_template_contract(api_client):
    admin = UserFactory(is_superuser=True, is_staff=True)
    assign_role(user=admin, role=ROLE_ADMIN)
    Category.objects.create(code="export-test", name="Nhóm xuất thử")
    api_client.force_authenticate(admin)

    response = api_client.get("/api/v1/exports/categories/")

    assert response.status_code == 200
    workbook = load_workbook(io.BytesIO(response.content))
    assert workbook["Data"]["A1"].value == "Mã (*)"
    assert workbook["Data"]["A2"].value == "code"
    assert workbook["Data"].row_dimensions[2].hidden is True
    parsed = parse_workbook(module="categories", content=response.content)
    assert parsed.issues == []
    assert any(row["code"] == "export-test" for row in parsed.rows)


@pytest.mark.django_db
def test_template_has_required_headers_and_database_dropdowns():
    Category.objects.get_or_create(code="coffee", defaults={"name": "Cà phê"})
    workbook = load_workbook(io.BytesIO(build_template("products")))

    assert workbook["Hướng dẫn"]["B1"].value == "2"
    assert workbook["Data"]["A1"].value == "Mã (*)"
    assert workbook["Data"]["C1"].value == "Danh mục (*)"
    assert workbook["Data"].row_dimensions[2].hidden is True
    assert any(
        validation.formula1 == "=category_codes"
        for validation in workbook["Data"].data_validations.dataValidation
    )

    Unit.objects.get_or_create(code="ml", defaults={"name": "Mililít", "symbol": "ml"})
    ingredient_workbook = load_workbook(io.BytesIO(build_template("ingredients")))
    assert any(
        validation.formula1 == "=unit_codes"
        for validation in ingredient_workbook["Data"].data_validations.dataValidation
    )


@pytest.mark.django_db
def test_create_mode_rejects_existing_code_and_provides_error_workbook(api_client):
    admin = UserFactory(is_superuser=True, is_staff=True)
    assign_role(user=admin, role=ROLE_ADMIN)
    Category.objects.create(code="new-coffee", name="Đã có")
    api_client.force_authenticate(admin)
    uploaded = SimpleUploadedFile(
        "categories.xlsx",
        category_file(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    preview = api_client.post(
        "/api/v1/imports/",
        {"module": "categories", "mode": "create", "file": uploaded},
        format="multipart",
    )
    assert preview.status_code == 201
    assert preview.data["status"] == "invalid"
    assert preview.data["error_rows"] == 1
    annotated = api_client.get(f"/api/v1/imports/{preview.data['id']}/error-file/")
    assert annotated.status_code == 200
    workbook = load_workbook(io.BytesIO(annotated.content))
    assert workbook["Errors"]["C2"].value == "already_exists"
