import hashlib
import io
from dataclasses import dataclass
from decimal import Decimal

from django.db import transaction
from django.utils import timezone
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from rest_framework.exceptions import ValidationError

from apps.catalog.models import Category, Product, Variant
from apps.inventory.models import Ingredient, Supplier, Unit

from .models import ImportJob, RowIssue

MAX_FILE_SIZE = 5 * 1024 * 1024
MAX_ROWS = 5_000
SUPPORTED_MODULES = {"categories", "products", "ingredients", "suppliers"}
MODULE_HEADERS = {
    "categories": ["code", "name", "description", "display_order", "is_active"],
    "products": [
        "code",
        "slug",
        "category_code",
        "name",
        "description",
        "price",
        "image_url",
        "is_featured",
        "is_active",
    ],
    "ingredients": ["code", "name", "base_unit_code", "track_lot", "minimum_level", "is_active"],
    "suppliers": ["code", "name", "contact_name", "phone", "email", "address", "is_active"],
}
FIELD_LABELS = {
    "code": "Mã",
    "name": "Tên",
    "description": "Mô tả",
    "display_order": "Thứ tự hiển thị",
    "is_active": "Đang sử dụng",
    "slug": "Đường dẫn",
    "category_code": "Danh mục",
    "price": "Giá bán",
    "image_url": "Đường dẫn ảnh",
    "is_featured": "Món nổi bật",
    "base_unit_code": "Đơn vị cơ sở",
    "track_lot": "Theo dõi theo lô",
    "minimum_level": "Mức tồn tối thiểu",
    "contact_name": "Người liên hệ",
    "phone": "Điện thoại",
    "email": "Email",
    "address": "Địa chỉ",
}
REQUIRED_FIELDS = {
    "categories": {"code", "name"},
    "products": {"code", "category_code", "name", "price"},
    "ingredients": {"code", "name", "base_unit_code"},
    "suppliers": {"code", "name"},
}
BOOLEAN_FIELDS = {"is_active", "is_featured", "track_lot"}
TEMPLATE_VERSION = "2"


@dataclass
class ParsedWorkbook:
    rows: list[dict]
    issues: list[dict]


def _safe_excel_value(value):
    if isinstance(value, bool):
        return "Có" if value else "Không"
    if not isinstance(value, str):
        return value
    return f"'{value}" if value.startswith(("=", "+", "-", "@")) else value


def build_template(module: str) -> bytes:
    _validate_module(module)
    workbook = Workbook()
    instructions = workbook.active
    instructions.title = "Hướng dẫn"
    instructions.append(["template_version", TEMPLATE_VERSION])
    instructions.append(["module", module])
    instructions.append(["Cách nhập", "Nhập dữ liệu từ dòng 3 của sheet Data."])
    instructions.append(["Trường bắt buộc", "Header màu nâu có dấu (*) phải được nhập."])
    instructions.append(["Danh sách chọn", "Ô màu xanh có danh sách xổ xuống từ dữ liệu hệ thống."])
    instructions.append(["Cập nhật", "Dùng __CLEAR__ để xóa trường không bắt buộc."])
    instructions.column_dimensions["A"].width = 22
    instructions.column_dimensions["B"].width = 72
    instructions.freeze_panes = "A3"
    for cell in instructions[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="6F3F25")
    data = workbook.create_sheet("Data")
    headers = MODULE_HEADERS[module]
    required = REQUIRED_FIELDS[module]
    data.append(
        [f"{FIELD_LABELS[key]} (*)" if key in required else FIELD_LABELS[key] for key in headers]
    )
    data.append(headers)
    data.row_dimensions[2].hidden = True
    data.freeze_panes = "A3"
    data.sheet_view.showGridLines = False
    thin = Side(style="thin", color="D8C8B5")
    for index, (cell, key) in enumerate(zip(data[1], headers, strict=False), start=1):
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="7A3F22" if key in required else "9B765E")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=thin)
        cell.comment = Comment(
            f"Trường {'bắt buộc' if key in required else 'không bắt buộc'} · mã kỹ thuật: {key}",
            "Kim Thành Coffee",
        )
        data.column_dimensions[get_column_letter(index)].width = max(16, len(cell.value) + 4)
    data.row_dimensions[1].height = 34
    _add_template_validations(workbook, data, module, headers)
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def _add_template_validations(workbook, data, module: str, headers: list[str]) -> None:
    reference = workbook.create_sheet("Danh sách chọn")
    reference.sheet_state = "hidden"
    reference_values: list[tuple[str, list[tuple[str, str]]]] = []
    if module == "products":
        reference_values.append(
            (
                "category_codes",
                list(
                    Category.objects.filter(is_active=True)
                    .order_by("name")
                    .values_list("code", "name")
                ),
            )
        )
    if module == "ingredients":
        reference_values.append(
            (
                "unit_codes",
                list(
                    Unit.objects.filter(is_active=True).order_by("name").values_list("code", "name")
                ),
            )
        )
    available_reference_names: set[str] = set()
    for column_index, (range_name, values) in enumerate(reference_values, start=1):
        reference.cell(1, column_index, "Mã")
        reference.cell(1, column_index + 1, "Tên hiển thị")
        for row_index, (code, name) in enumerate(values, start=2):
            reference.cell(row_index, column_index, code)
            reference.cell(row_index, column_index + 1, name)
        if values:
            column = get_column_letter(column_index)
            workbook.defined_names.add(
                DefinedName(
                    range_name,
                    attr_text=f"'Danh sách chọn'!${column}$2:${column}${len(values) + 1}",
                )
            )
            available_reference_names.add(range_name)
    reference_by_field = {"category_code": "category_codes", "base_unit_code": "unit_codes"}
    for index, key in enumerate(headers, start=1):
        column = get_column_letter(index)
        validation = None
        if reference_by_field.get(key) in available_reference_names:
            validation = DataValidation(type="list", formula1=f"={reference_by_field[key]}")
        elif key in BOOLEAN_FIELDS:
            validation = DataValidation(type="list", formula1='"Có,Không"')
        if validation:
            validation.error = "Vui lòng chọn một giá trị trong danh sách."
            validation.errorTitle = "Giá trị không hợp lệ"
            validation.prompt = "Chọn giá trị từ danh sách xổ xuống."
            validation.promptTitle = "Danh sách chọn"
            validation.showErrorMessage = True
            validation.showInputMessage = True
            data.add_data_validation(validation)
            validation.add(f"{column}3:{column}{MAX_ROWS + 2}")


def _validate_module(module: str) -> None:
    if module not in SUPPORTED_MODULES:
        raise ValidationError({"module": "Module import chưa được hỗ trợ."})


def _dependency_fingerprint(module: str) -> str:
    if module == "products":
        values = "|".join(Category.objects.order_by("code").values_list("code", flat=True))
    elif module == "ingredients":
        values = "|".join(Unit.objects.order_by("code").values_list("code", flat=True))
    else:
        values = module
    return hashlib.sha256(values.encode()).hexdigest()


def parse_workbook(*, module: str, content: bytes, mode: str = "upsert") -> ParsedWorkbook:
    _validate_module(module)
    if len(content) > MAX_FILE_SIZE:
        raise ValidationError({"file": "File vượt giới hạn 5 MB."})
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
    except Exception as exc:
        raise ValidationError({"file": "File XLSX không hợp lệ."}) from exc
    if "Data" not in workbook.sheetnames:
        raise ValidationError({"file": "Thiếu sheet Data."})
    instruction_sheet = "Hướng dẫn" if "Hướng dẫn" in workbook.sheetnames else "Instructions"
    if instruction_sheet not in workbook.sheetnames:
        raise ValidationError({"file": "Phiên bản template không được hỗ trợ."})
    version = str(workbook[instruction_sheet]["B1"].value or "")
    if version not in {"1", TEMPLATE_VERSION}:
        raise ValidationError({"file": "Phiên bản template không được hỗ trợ."})
    sheet = workbook["Data"]
    header_row = 2 if version == TEMPLATE_VERSION else 1
    header = [
        str(value or "").strip()
        for value in next(sheet.iter_rows(min_row=header_row, max_row=header_row, values_only=True))
    ]
    expected = MODULE_HEADERS[module]
    if header != expected:
        raise ValidationError({"file": "Header không đúng template hiện hành."})
    rows: list[dict] = []
    issues: list[dict] = []
    seen: set[str] = set()
    existing_codes = _existing_codes(module)
    first_data_row = header_row + 1
    for row_number, values in enumerate(
        sheet.iter_rows(min_row=first_data_row, values_only=True), start=first_data_row
    ):
        if row_number > MAX_ROWS + first_data_row - 1:
            raise ValidationError({"file": "File vượt giới hạn 5.000 dòng."})
        if all(value in (None, "") for value in values):
            continue
        row = dict(zip(header, values, strict=False))
        code = str(row.get("code") or "").strip()
        if not code:
            issues.append(_issue(row_number, "code", "required", "Code là bắt buộc.", ""))
        elif code in seen:
            issues.append(
                _issue(row_number, "code", "duplicate", "Code bị trùng trong file.", code)
            )
        seen.add(code)
        if code and mode == ImportJob.Mode.CREATE and code in existing_codes:
            issues.append(_issue(row_number, "code", "already_exists", "Code đã tồn tại.", code))
        if code and mode == ImportJob.Mode.UPDATE and code not in existing_codes:
            issues.append(_issue(row_number, "code", "not_found", "Code chưa tồn tại.", code))
        if not str(row.get("name") or "").strip():
            issues.append(_issue(row_number, "name", "required", "Tên là bắt buộc.", ""))
        for field in REQUIRED_FIELDS[module] - {"code", "name"}:
            if row.get(field) in (None, ""):
                issues.append(
                    _issue(
                        row_number,
                        field,
                        "required",
                        f"{FIELD_LABELS[field]} là bắt buộc.",
                        "",
                    )
                )
        if (
            module == "products"
            and not Category.objects.filter(
                code=str(row.get("category_code") or "").strip()
            ).exists()
        ):
            issues.append(
                _issue(
                    row_number,
                    "category_code",
                    "missing_reference",
                    "Không tìm thấy danh mục.",
                    row.get("category_code"),
                )
            )
        if (
            module == "ingredients"
            and not Unit.objects.filter(code=str(row.get("base_unit_code") or "").strip()).exists()
        ):
            issues.append(
                _issue(
                    row_number,
                    "base_unit_code",
                    "missing_reference",
                    "Không tìm thấy đơn vị.",
                    row.get("base_unit_code"),
                )
            )
        rows.append({key: value for key, value in row.items()})
    return ParsedWorkbook(rows=rows, issues=issues)


def _issue(row: int, column: str, code: str, message: str, raw_value) -> dict:
    return {
        "row_number": row,
        "column": column,
        "field_code": column,
        "error_code": code,
        "message": message,
        "raw_value": str(raw_value or "")[:255],
    }


@transaction.atomic
def create_import_job(*, actor, module: str, mode: str, uploaded_file) -> ImportJob:
    content = uploaded_file.read()
    parsed = parse_workbook(module=module, content=content, mode=mode)
    error_rows = len({issue["row_number"] for issue in parsed.issues})
    existing_codes = _existing_codes(module)
    new_rows = sum(1 for row in parsed.rows if str(row.get("code")) not in existing_codes)
    update_rows = len(parsed.rows) - new_rows
    job = ImportJob.objects.create(
        module=module,
        mode=mode,
        status=ImportJob.Status.INVALID if parsed.issues else ImportJob.Status.READY,
        original_name=uploaded_file.name,
        file_hash=hashlib.sha256(content).hexdigest(),
        file_content=content,
        total_rows=len(parsed.rows),
        new_rows=new_rows,
        update_rows=update_rows,
        error_rows=error_rows,
        preview=parsed.rows[:100],
        dependency_fingerprint=_dependency_fingerprint(module),
        created_by=actor,
        validated_at=timezone.now(),
    )
    RowIssue.objects.bulk_create([RowIssue(job=job, **issue) for issue in parsed.issues])
    return job


def _existing_codes(module: str) -> set[str]:
    model = {
        "categories": Category,
        "products": Product,
        "ingredients": Ingredient,
        "suppliers": Supplier,
    }[module]
    return set(model.objects.values_list("code", flat=True))


def _as_bool(value, default=True) -> bool:
    if value in (None, ""):
        return default
    return str(value).strip().lower() in {"1", "true", "yes", "có", "x"}


@transaction.atomic
def confirm_import(*, job: ImportJob, actor, idempotency_key: str) -> ImportJob:
    existing = ImportJob.objects.filter(
        created_by=actor, idempotency_key=idempotency_key, status=ImportJob.Status.CONFIRMED
    ).first()
    if existing:
        return existing
    job = ImportJob.objects.select_for_update().get(pk=job.pk)
    if job.status == ImportJob.Status.CONFIRMED:
        if job.idempotency_key == idempotency_key:
            return job
        raise ValidationError("Import đã được xác nhận.", code="already_confirmed")
    if job.status != ImportJob.Status.READY:
        raise ValidationError("Import chưa sẵn sàng để xác nhận.")
    if job.created_by_id != actor.id:
        raise ValidationError("Bạn không thể xác nhận import của người khác.")
    if job.dependency_fingerprint != _dependency_fingerprint(job.module):
        job.status = ImportJob.Status.STALE
        job.save(update_fields=["status"])
        raise ValidationError(
            "Dữ liệu tham chiếu đã thay đổi. Vui lòng tải lại file.", code="stale"
        )
    parsed = parse_workbook(module=job.module, content=bytes(job.file_content), mode=job.mode)
    if parsed.issues:
        raise ValidationError("File không còn hợp lệ.")
    for row in parsed.rows:
        _write_row(job.module, job.mode, row)
    job.status = ImportJob.Status.CONFIRMED
    job.idempotency_key = idempotency_key
    job.confirmed_at = timezone.now()
    job.save(update_fields=["status", "idempotency_key", "confirmed_at"])
    return job


def _write_row(module: str, mode: str, row: dict) -> None:
    code = str(row["code"]).strip()
    if module == "categories":
        values = {
            "name": str(row["name"]).strip(),
            "description": _nullable_text(row.get("description")),
            "display_order": int(row.get("display_order") or 0),
            "is_active": _as_bool(row.get("is_active")),
        }
        Category.objects.update_or_create(code=code, defaults=values)
    elif module == "products":
        category = Category.objects.get(code=str(row["category_code"]).strip())
        values = {
            "slug": str(row.get("slug") or code).strip(),
            "category": category,
            "name": str(row["name"]).strip(),
            "description": _nullable_text(row.get("description")),
            "image_url": _nullable_text(row.get("image_url")),
            "is_featured": _as_bool(row.get("is_featured"), False),
            "is_active": _as_bool(row.get("is_active")),
        }
        product, _ = Product.objects.update_or_create(code=code, defaults=values)
        price = Decimal(str(row.get("price") or "0"))
        Variant.objects.update_or_create(
            code=f"{code}-default",
            defaults={"product": product, "name": "Mặc định", "price": price, "is_default": True},
        )
    elif module == "ingredients":
        unit = Unit.objects.get(code=str(row["base_unit_code"]).strip())
        Ingredient.objects.update_or_create(
            code=code,
            defaults={
                "name": str(row["name"]).strip(),
                "base_unit": unit,
                "track_lot": _as_bool(row.get("track_lot"), False),
                "minimum_level": Decimal(str(row.get("minimum_level") or "0")),
                "is_active": _as_bool(row.get("is_active")),
            },
        )
    else:
        Supplier.objects.update_or_create(
            code=code,
            defaults={
                "name": str(row["name"]).strip(),
                "contact_name": _nullable_text(row.get("contact_name")),
                "phone": _nullable_text(row.get("phone")),
                "email": _nullable_text(row.get("email")),
                "address": _nullable_text(row.get("address")),
                "is_active": _as_bool(row.get("is_active")),
            },
        )


def export_module(module: str) -> bytes:
    _validate_module(module)
    workbook = load_workbook(io.BytesIO(build_template(module)))
    sheet = workbook["Data"]
    headers = MODULE_HEADERS[module]
    for row_index, row in enumerate(_export_rows(module), start=3):
        for column_index, header in enumerate(headers, start=1):
            sheet.cell(row_index, column_index, _safe_excel_value(row.get(header)))
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def _nullable_text(value) -> str:
    return "" if value == "__CLEAR__" else str(value or "")


def build_error_workbook(job: ImportJob) -> bytes:
    workbook = load_workbook(io.BytesIO(bytes(job.file_content)))
    if "Errors" in workbook.sheetnames:
        del workbook["Errors"]
    sheet = workbook.create_sheet("Errors")
    sheet.append(["row_number", "column", "error_code", "message", "raw_value"])
    for issue in job.issues.all():
        sheet.append(
            [issue.row_number, issue.column, issue.error_code, issue.message, issue.raw_value]
        )
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def _export_rows(module: str) -> list[dict]:
    if module == "categories":
        return list(Category.objects.values(*MODULE_HEADERS[module]))
    if module == "suppliers":
        return list(Supplier.objects.values(*MODULE_HEADERS[module]))
    if module == "ingredients":
        return [
            {
                "code": item.code,
                "name": item.name,
                "base_unit_code": item.base_unit.code,
                "track_lot": item.track_lot,
                "minimum_level": item.minimum_level,
                "is_active": item.is_active,
            }
            for item in Ingredient.objects.select_related("base_unit")
        ]
    return [
        {
            "code": item.code,
            "slug": item.slug,
            "category_code": item.category.code,
            "name": item.name,
            "description": item.description,
            "price": item.variants.order_by("display_order").values_list("price", flat=True).first()
            or 0,
            "image_url": item.image_url,
            "is_featured": item.is_featured,
            "is_active": item.is_active,
        }
        for item in Product.objects.select_related("category").prefetch_related("variants")
    ]
